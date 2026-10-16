import json

import pytest

pytest.importorskip("matplotlib")
pytest.importorskip("pptx")  # optional: pip install -e ".[dev]"

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

from screening.viz.figures import BUILDERS, load_payload, make_figures  # noqa: E402
from scripts.build_deck import build_deck  # noqa: E402


def _payload():
    def rate(group, attendance, n=100):
        return {"group": group, "n": n, "attendance": attendance,
                "ci_low": attendance - 0.01, "ci_high": attendance + 0.01}
    segments = [
        {"segment": "first_invitation", "n": 600, "share": 0.6,
         "non_attendance": 0.4, "mean_predicted_risk": 0.4},
        {"segment": "repeated_nonresponder", "n": 400, "share": 0.4,
         "non_attendance": 0.9, "mean_predicted_risk": 0.9},
    ]
    return {
        "meta": {"n_invitations": 1000, "n_clients": 800, "years": [2020, 2024],
                 "attendance_overall": 0.2, "test_year": 2024, "n_test": 200,
                 "min_cell": 20},
        "test_metrics": [{"metric": "roc_auc", "model": 0.84, "baseline": 0.78}],
        "attendance_by": {
            "prior_invitations": [rate(0, 0.51), rate(1, 0.30), rate(2, 0.03)],
            "months_since_last_screening": [rate("0", 0.02), rate("1", 0.02),
                                             rate("2-6", 0.03), rate("7-24", 0.4),
                                             rate("none_in_5y", 0.1)],
        },
        "deciles": [{"decile": i, "attendance": 0.7 - 0.065 * (i - 1), "n": 20}
                    for i in range(1, 11)],
        "targeting": [{"share": 0.1, "n": 200, "top_wasted_model": 5,
                       "top_wasted_rule": 9, "bottom_attendance_model": 0.68,
                       "bottom_attendance_rule": 0.51}],
        "feature_importance": [
            {"feature": f"feature_{i}", "importance": 1 / i,
             "direction": 1 if i % 2 else -1} for i in range(1, 11)
        ],
        "segments": segments,
        "segment_response_speed": [
            {"segment": row["segment"], "3m": 0.05, "6m": 0.1, "12m": 0.2}
            for row in segments
        ],
        "experiment_design": [
            {"segment": row["segment"], "baseline_attendance": 1 - row["non_attendance"],
             "mde_pp": 2, "n_per_arm": 100 + index * 50,
             "segment_size": row["n"], "feasible": True}
            for index, row in enumerate(segments)
        ],
        "recently_screened_per_year": [rate(2023, 0.02, 70), rate(2024, 0.02, 80)],
        "interventions": {"first_invitation": "Standard letter"},
        "calibration": [],
        "fairness": {},
    }


def test_pitch_figures_and_deck_are_reproducible_from_json(tmp_path):
    payload = _payload()
    json_path = tmp_path / "data.json"
    json_path.write_text(json.dumps(payload), encoding="utf-8")

    figures = tmp_path / "figures"
    outputs = make_figures(load_payload(json_path), figures)
    assert len(outputs) == len(BUILDERS) == 8
    assert all(Image.open(path).size == (1600, 900) for path in outputs)

    deck = build_deck(payload, figures, tmp_path / "deck.pptx")
    presentation = Presentation(deck)
    assert len(presentation.slides) == 12
    slide_text = "\n".join(shape.text for shape in presentation.slides[4].shapes if hasattr(shape, "text"))
    assert "0.840" in slide_text and "68%" in slide_text and "51%" in slide_text


def test_missing_export_key_creates_explicit_todo(tmp_path):
    payload = _payload()
    payload.pop("targeting")
    output = make_figures(payload, tmp_path)
    assert (tmp_path / "model_vs_rule.png") in output

